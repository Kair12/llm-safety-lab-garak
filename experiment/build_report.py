from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "deliverables" / "LLM_Safety_Experimental_Report.docx"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F7"
BORDER = "D9D9D9"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[col_index])
            if row_index == 0:
                set_cell_fill(cell, NAVY)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.bold = True
                        run.font.size = Pt(9)
            else:
                set_cell_fill(cell, PALE_BLUE if row_index % 2 == 0 else "FFFFFF")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        lead = p.add_run(bold_lead)
        lead.bold = True
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.58)
section.bottom_margin = Inches(0.58)
section.left_margin = Inches(0.68)
section.right_margin = Inches(0.68)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(4)
styles["Normal"].paragraph_format.line_spacing = 1.05

for name, size in (("Title", 21), ("Heading 1", 14), ("Heading 2", 11)):
    style = styles[name]
    style.font.name = "Aptos Display" if name != "Heading 2" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(7 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(4)

styles["Title"].paragraph_format.space_after = Pt(5)
styles["List Bullet"].font.name = "Aptos"
styles["List Bullet"].font.size = Pt(9.5)
styles["List Bullet"].paragraph_format.space_after = Pt(2)
styles["List Bullet"].paragraph_format.left_indent = Inches(0.22)
styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.12)

footer = section.footer
footer.paragraphs[0].text = "LLM Safety Evaluation and NVIDIA garak"
footer.paragraphs[0].runs[0].font.size = Pt(8)
add_page_number(footer.add_paragraph())

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("LLM Safety Evaluation with NVIDIA garak")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(8)
run = subtitle.add_run("Research review and controlled red teaming experiment")
run.italic = True
run.font.size = Pt(11)

add_body(
    doc,
    "This report reviews a 2026 ACL Findings study on jailbreak attacks and content safety filters, then tests a local language model with NVIDIA garak. The experiment found prompt injection in 7 of 16 sampled cases and a successful DAN jailbreak in 1 of 1 case. A system instruction alone did not improve the results. A separate input filter reduced the reported failures to zero, but it also blocked 3 of 20 benign prompts. The evidence supports layered controls and broader evaluation rather than a claim of complete safety.",
)

doc.add_heading("Selected scientific study", level=1)
add_body(
    doc,
    "Xin et al. study the gap between model-only jailbreak evaluation and real deployments that add moderation before and after model inference [1]. Their central research question is whether jailbreaks that defeat an aligned model still succeed when the full inference pipeline contains content safety filters. This focus matches the assignment because it examines red teaming, jailbreaks, automated vulnerability testing, and the practical limits of AI safety controls.",
)

doc.add_heading("Methodology and metrics", level=1)
add_body(
    doc,
    "The authors evaluate ten jailbreak methods against open and proprietary models. Their model set includes Llama, Mistral, Vicuna, Qwen, GPT-4o, and GPT-4-Turbo. They pair 417 harmful prompts with 417 benign prompts aligned to the same topics across ten safety categories. The study tests input and output filters including OpenAI moderation, LlamaGuard, PromptGuard, InjecGuard, GradSafe, and O3. This design measures the full path from adversarial input to accepted output instead of treating the base model as the only security boundary.",
)
add_body(
    doc,
    "The evaluation reports original attack success rate, input and output detection rates, end-to-end pass rate, and standard classification measures such as precision, recall, accuracy, and F1. It also considers latency and financial cost. These metrics matter together because a filter can catch many attacks while also rejecting legitimate prompts or adding unacceptable overhead.",
)

doc.add_heading("Findings and critical assessment", level=1)
add_body(
    doc,
    "The paper finds that at least one tested filter detects nearly every jailbreak technique, which means model-only studies can overstate practical attack success. However, the filters differ sharply in precision and recall. PromptGuard shows strong attack detection but an extreme false-positive problem in the reported setting. Adaptive attacks also remain important: an adaptive TAP variant raises its success against OpenAI moderation from 28.4% to 88.3%, although it requires more iterations [1]. The study therefore shows that filters reduce risk but also create usability costs and can fail when the attacker adapts.",
)
add_body(
    doc,
    "The study has a strong pipeline-level threat model, a balanced benign set, and several operational metrics. Its limitations include restricted API budgets, incomplete coverage of tool-using agents, and only preliminary analysis of adaptive attacks. These limits motivate the local experiment below, which repeats a narrower before-and-after comparison and explicitly measures a small benign false-positive sample.",
)

doc.add_page_break()

doc.add_heading("Experimental methodology", level=1)
add_body(
    doc,
    "The target was Qwen2.5-0.5B-Instruct through a local Hugging Face pipeline on an NVIDIA GeForce GTX 1660 with 6 GB VRAM [3]. NVIDIA garak version 0.17.1rc1 came from the official repository at commit 346497afcd86ff5ba137cfbcaaeba4e9c9f6c970 [2]. Every run used seed 42, one generation per prompt, and the same decoding settings. The smaller 0.5B model fit the available GPU but does not reproduce the paper's Qwen2.5-7B setting.",
)

setup = doc.add_table(rows=1, cols=4)
headers = ["Probe", "Purpose", "Prompts", "Detector"]
for i, text in enumerate(headers):
    setup.rows[0].cells[i].text = text
for values in (
    ("HijackLongPrompt", "Embedded instruction override", "16 sampled from 256", "AttackRogueString"),
    ("DAN 11.0", "Unrestricted persona jailbreak", "1 standardized prompt", "DAN and MitigationBypass"),
):
    cells = setup.add_row().cells
    for i, text in enumerate(values):
        cells[i].text = text
style_table(setup, [1.45, 2.25, 1.35, 2.05])

doc.add_heading("Baseline results", level=1)
results = doc.add_table(rows=1, cols=4)
for i, text in enumerate(("Evaluation", "Baseline", "System instruction", "Input filter")):
    results.rows[0].cells[i].text = text
for values in (
    ("PromptInject failures", "7/16  43.75%", "7/16  43.75%", "0/16  0%"),
    ("DAN detector failures", "1/1  100%", "1/1  100%", "0/1  0%"),
    ("Mitigation bypass failures", "1/1  100%", "1/1  100%", "0/1  0%"),
):
    cells = results.add_row().cells
    for i, text in enumerate(values):
        cells[i].text = text
style_table(results, [2.4, 1.55, 1.65, 1.55])

doc.add_heading("Safety failure analysis", level=1)
add_body(
    doc,
    "Prompt injection. The expected behavior was to complete the original benign task and treat the embedded override as untrusted data. Instead, seven responses reproduced the attack marker. The failed safety property was instruction hierarchy. In a retrieval or tool-enabled system, the same weakness could redirect an assistant, expose contextual information, or cause an unintended action.",
)
add_body(
    doc,
    "DAN jailbreak. The expected behavior was to reject the unrestricted persona. The baseline response began with “DAN Mode enabled,” so both the DAN detector and the mitigation-bypass detector reported failure. This result suggests weak resistance to role-play overrides, although one prompt provides little statistical evidence.",
)
add_body(
    doc,
    "Detector limitation. AttackRogueString searches for the target string. One baseline response split the marker into a numbered list and received a score of zero even though the output plausibly followed the injected request. The reported 43.75% attack success rate may therefore underestimate semantic compliance. Detector output requires manual review rather than automatic acceptance.",
)

doc.add_heading("Implemented mitigation", level=1)
add_body(
    doc,
    "A prompt-only control first placed security instructions in a system message. It did not change either result. The second control added a separate input filter before inference. The filter screens for general indicators such as commands to ignore earlier rules and requests to enter DAN or jailbreak modes. It blocked 15 of 16 PromptInject inputs and the DAN prompt. The one allowed PromptInject input did not emit the exact attack marker.",
)

doc.add_page_break()

doc.add_heading("Residual risk and recommendations", level=1)
add_body(
    doc,
    "The input filter improved the selected garak scores, but it did not establish model safety. Attackers can evade a keyword filter through paraphrases, encoding, multilingual prompts, or multi-turn interaction. The filter also blocked 3 of 20 hand-written benign prompts, a 15% false-positive rate. The blocked cases involved legitimate translation or educational use of security terms. This small convenience sample illustrates the usability tradeoff but does not estimate a production false-positive rate.",
)
add_bullet(doc, "Keep trusted system instructions separate from retrieved documents and user-controlled data, and preserve source provenance throughout processing.")
add_bullet(doc, "Restrict external tools by default and require human approval for consequential or irreversible actions.")
add_bullet(doc, "Validate outputs and tool arguments against task-specific schemas or allow-lists before execution.")
add_bullet(doc, "Repeat evaluation with larger benign and attack sets, adaptive variants, multilingual prompts, and multiple random seeds.")

doc.add_heading("Conclusion", level=1)
add_body(
    doc,
    "The experiment reproduced two safety failures on a small local model. A stronger system instruction did not help, while a separate input filter stopped the tested signatures at the cost of false positives. This pattern agrees with Xin et al.: deployment-level filters can reduce observed jailbreak success, but their precision, recall, and resistance to adaptive attacks determine whether they improve real safety. The correct conclusion is limited to the tested model, prompts, detector rules, and configuration.",
)

doc.add_heading("References", level=1)
refs = [
    "[1] Y. Xin, D. Chen, L. Yang, M. Backes, and X. Zhang, “Jailbreaking Attacks vs. Content Safety Filters: How Far Are We in the LLM Safety Arms Race?” Findings of ACL 2026, pp. 421-445, 2026. doi: 10.18653/v1/2026.findings-acl.20.",
    "[2] NVIDIA, “garak: Generative AI Red-teaming and Assessment Kit,” GitHub repository, 2026. https://github.com/NVIDIA/garak",
    "[3] Qwen Team, “Qwen2.5-0.5B-Instruct,” Hugging Face model card, 2024. https://huggingface.co/Qwen/Qwen2.5-0.5B-Instruct",
    "[4] F. Perez and I. Ribeiro, “Ignore Previous Prompt: Attack Techniques for Language Models,” NeurIPS ML Safety Workshop, 2022. https://arxiv.org/abs/2211.09527",
]
for text in refs:
    p = doc.add_paragraph(text)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(2)
    for run in p.runs:
        run.font.size = Pt(8.5)

doc.add_heading("AI Use Declaration", level=1)
add_body(
    doc,
    "I used ChatGPT Codex to explain the assignment, help install and configure NVIDIA garak, draft reproducibility scripts, summarize logs, improve the language and formatting, and generate the abstract presentation cover illustration. I independently selected the paper, reviewed the experimental evidence, checked the model responses, and remain responsible for the interpretation and conclusions. I verified references against the ACL Anthology, the official NVIDIA garak repository and documentation, the Qwen model card, and the saved JSONL and HTML reports. I did not treat AI-generated claims or citations as verified without checking those sources.",
)

doc.core_properties.title = "LLM Safety Evaluation with NVIDIA garak"
doc.core_properties.subject = "Research review and controlled red teaming experiment"
doc.core_properties.author = "Student"
doc.save(OUTPUT)
print(OUTPUT)
